"""เพิ่มกราฟความเร็ว-เวลาในชีต 'ข้อมูลดิบ' ของ movement_report.xlsx ที่มีอยู่

กราฟใช้ความเร็วเฉลี่ยและความเร็วสูงสุดของทุกเป้า ณ เวลาเดียวกัน เพื่อไม่ให้
การเรียงแถวของ ArUco คนละ ID ทำให้เส้นกราฟกระโดดผิดความหมาย
"""

from __future__ import annotations

import argparse
import shutil
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


RAW_SHEET_NAME = "ข้อมูลดิบ"
TIME_HEADER = "เวลาสะสม (วินาที)"
SPEED_HEADER = "ความเร็ว (มม./วินาที)"
SUMMARY_HEADERS = [
    "เวลาสะสม (วินาที)",
    "ความเร็วเฉลี่ย (มม./วินาที)",
    "ความเร็วสูงสุด (มม./วินาที)",
]


def header_columns(sheet) -> dict[str, int]:
    columns: dict[str, int] = {}
    for cell in sheet[1]:
        if isinstance(cell.value, str) and cell.value and cell.value not in columns:
            columns[cell.value] = cell.column
    return columns


def add_speed_chart(report_path: Path) -> None:
    workbook = load_workbook(report_path)
    if RAW_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"ไม่พบชีต '{RAW_SHEET_NAME}' ใน {report_path}")
    sheet = workbook[RAW_SHEET_NAME]
    columns = header_columns(sheet)
    if TIME_HEADER not in columns or SPEED_HEADER not in columns:
        raise ValueError("ไม่พบคอลัมน์เวลาสะสมหรือความเร็วในข้อมูลดิบ")

    grouped: dict[float, list[float]] = defaultdict(list)
    for row in range(2, sheet.max_row + 1):
        elapsed = sheet.cell(row=row, column=columns[TIME_HEADER]).value
        speed = sheet.cell(row=row, column=columns[SPEED_HEADER]).value
        if isinstance(elapsed, (int, float)) and isinstance(speed, (int, float)):
            grouped[round(float(elapsed), 3)].append(float(speed))
    if not grouped:
        raise ValueError("ไม่พบข้อมูลความเร็วที่ใช้สร้างกราฟ")

    # P:R is deliberately separate from the original raw-data table A:O.
    # Reuse it on a re-run, rather than appending duplicate helper columns/charts.
    start_column = columns.get(SUMMARY_HEADERS[1], 17) - 1
    end_column = start_column + 2
    for column, header in enumerate(SUMMARY_HEADERS, start=start_column):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        sheet.column_dimensions[get_column_letter(column)].width = 24

    for row, elapsed in enumerate(sorted(grouped), start=2):
        speeds = grouped[elapsed]
        values = (elapsed, sum(speeds) / len(speeds), max(speeds))
        for column, value in enumerate(values, start=start_column):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.number_format = "0.000"

    # The raw sheet is reserved for this one speed chart. Other-sheet charts remain untouched.
    sheet._charts = []
    chart = LineChart()
    chart.title = "ความเร็วเฉลี่ยและสูงสุดเทียบกับเวลา"
    chart.style = 13
    chart.y_axis.title = "ความเร็ว (มม./วินาที)"
    chart.x_axis.title = "เวลาสะสม (วินาที)"
    chart.height = 12
    chart.width = 25
    data = Reference(sheet, min_col=start_column + 1, max_col=end_column,
                     min_row=1, max_row=1 + len(grouped))
    categories = Reference(sheet, min_col=start_column, min_row=2,
                           max_row=1 + len(grouped))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.series[0].tx = SeriesLabel(v=SUMMARY_HEADERS[1])
    chart.series[1].tx = SeriesLabel(v=SUMMARY_HEADERS[2])
    chart.series[0].graphicalProperties.line.solidFill = "1F4E78"
    chart.series[1].graphicalProperties.line.solidFill = "C00000"
    sheet.add_chart(chart, f"{get_column_letter(end_column + 2)}2")
    workbook.save(report_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("reports", nargs="+", type=Path, help="ไฟล์ movement_report.xlsx")
    args = parser.parse_args()
    for report_path in args.reports:
        if not report_path.is_file():
            raise FileNotFoundError(report_path)
        backup_path = report_path.with_name("movement_report_before_speed_chart.xlsx")
        if not backup_path.exists():
            shutil.copy2(report_path, backup_path)
        add_speed_chart(report_path)
        print(f"เพิ่มกราฟแล้ว: {report_path}")
        print(f"ข้อมูลสำรอง: {backup_path}")


if __name__ == "__main__":
    main()
