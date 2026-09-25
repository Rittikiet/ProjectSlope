"""Correct selected sessions after discovering the real ArUco marker side length.

Example:
    python correct_marker_size_sessions.py --apply --session-ids \
        20260909_163025_985761 20260909_164913_005410

The command creates a timestamped database backup before changing any records.
"""

import argparse
import csv
from datetime import datetime
from pathlib import Path
import shutil
import sqlite3

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RAW_HEADERS = [
    "วันเวลาที่บันทึก", "รหัส ArUco", "ตำแหน่ง X (เมตร)", "ตำแหน่ง Y (เมตร)",
    "ตำแหน่ง Z (เมตร)", "ไฟล์ภาพประกอบ", "จุดกึ่งกลาง X (พิกเซล)",
    "จุดกึ่งกลาง Y (พิกเซล)", "การเปลี่ยน X (มม.)", "การเปลี่ยน Y (มม.)",
    "การเปลี่ยน Z (มม.)", "ระยะการเคลื่อนที่ (มม.)", "เวลาสะสม (วินาที)",
    "ความเร็ว (มม./วินาที)", "มุมเอียง (องศา)",
]


def write_csv(path: Path, headers, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(headers)
        writer.writerows(rows)


def write_excel(path: Path, comparison_rows, marker_ids, raw_rows):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "กราฟและข้อมูล"
    headers = ["เวลาสะสม (วินาที)", "วันเวลาที่บันทึก",
               *[f"เป้า ID {marker_id} (มม.)" for marker_id in marker_ids]]
    summary.append(headers)
    for elapsed, recorded_at, distances in comparison_rows:
        summary.append([elapsed, recorded_at, *[distances.get(marker_id) for marker_id in marker_ids]])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    summary.freeze_panes = "C2"
    summary.auto_filter.ref = summary.dimensions
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 32
    for column in range(3, len(headers) + 1):
        summary.cell(1, column).column_letter
        summary.column_dimensions[summary.cell(1, column).column_letter].width = 18
    if comparison_rows and marker_ids:
        chart = LineChart()
        chart.title = "ระยะการเคลื่อนที่เทียบกับเวลา"
        chart.y_axis.title = "ระยะการเคลื่อนที่ (มม.)"
        chart.x_axis.title = "เวลาสะสม (วินาที)"
        chart.height = 14
        chart.width = 28
        chart.add_data(Reference(summary, min_col=3, max_col=2 + len(marker_ids), min_row=1,
                                 max_row=len(comparison_rows) + 1), titles_from_data=True)
        chart.set_categories(Reference(summary, min_col=1, min_row=2,
                                       max_row=len(comparison_rows) + 1))
        summary.add_chart(chart, f"A{len(comparison_rows) + 4}")

    raw = workbook.create_sheet("ข้อมูลดิบ")
    raw.append(RAW_HEADERS)
    raw_rows_for_excel = []
    for row in raw_rows:
        row = list(row)
        if row[5]:
            row[5] = f"snapshots/{row[5]}"
        raw.append(row)
        raw_rows_for_excel.append(row)
    for cell in raw[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    raw.freeze_panes = "A2"
    raw.auto_filter.ref = raw.dimensions
    for index, header in enumerate(RAW_HEADERS, start=1):
        raw.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 3, 15), 28)
    speed_by_time = {}
    for raw_row in raw_rows:
        elapsed_seconds, speed = raw_row[12], raw_row[13]
        if elapsed_seconds is None or speed is None:
            continue
        key = round(float(elapsed_seconds), 3)
        point = speed_by_time.setdefault(key, (float(key), []))
        point[1].append(float(speed))
    speed_summary = [
        (elapsed, sum(speeds) / len(speeds), max(speeds))
        for elapsed, speeds in sorted(speed_by_time.values())
    ]
    if speed_summary:
        summary_start = len(RAW_HEADERS) + 2
        headers = ["เวลาสะสม (วินาที)", "ความเร็วเฉลี่ย (มม./วินาที)",
                   "ความเร็วสูงสุด (มม./วินาที)"]
        for index, header in enumerate(headers, start=summary_start):
            cell = raw.cell(1, index, header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            raw.column_dimensions[get_column_letter(index)].width = 22
        for row_index, values in enumerate(speed_summary, start=2):
            for column_index, value in enumerate(values, start=summary_start):
                raw.cell(row_index, column_index, value).number_format = "0.000"
        chart = LineChart()
        chart.title = "ความเร็วเฉลี่ยและสูงสุดเทียบกับเวลา"
        chart.y_axis.title = "ความเร็ว (มม./วินาที)"
        chart.x_axis.title = "เวลาสะสม (วินาที)"
        chart.height = 14
        chart.width = 28
        chart.add_data(Reference(raw, min_col=summary_start + 1, max_col=summary_start + 2,
                                 min_row=1, max_row=len(speed_summary) + 1), titles_from_data=True)
        chart.set_categories(Reference(raw, min_col=summary_start, min_row=2,
                                       max_row=len(speed_summary) + 1))
        chart.series[0].tx = SeriesLabel(v=headers[1])
        chart.series[1].tx = SeriesLabel(v=headers[2])
        chart.series[0].graphicalProperties.line.solidFill = "1F4E78"
        chart.series[1].graphicalProperties.line.solidFill = "C00000"
        raw.add_chart(chart, f"{get_column_letter(summary_start + 4)}2")
    workbook.save(path)


def regenerate_reports(connection, session_id: str, sessions_directory: Path):
    rows = connection.execute(
        """
        SELECT recorded_at, marker_id, x_m, y_m, z_m, image_file,
               center_x_px, center_y_px, dx_mm, dy_mm, dz_mm, distance_mm,
               elapsed_seconds, speed_mm_s, tilt_deg
        FROM measurements WHERE session_id = ? ORDER BY elapsed_seconds, marker_id
        """,
        (session_id,),
    ).fetchall()
    if not rows:
        raise ValueError(f"ไม่พบข้อมูล session: {session_id}")
    output = sessions_directory / f"session_{session_id}"
    output.mkdir(parents=True, exist_ok=True)

    raw_rows = []
    points = {}
    marker_ids = set()
    for row in rows:
        raw_row = list(row)
        if raw_row[5]:
            raw_row[5] = f"snapshots/{raw_row[5]}"
        raw_rows.append(raw_row)
        recorded_at, marker_id = row[0], row[1]
        elapsed, distance = row[12], row[11]
        marker_ids.add(marker_id)
        point = points.setdefault(round(elapsed, 3), (elapsed, recorded_at, {}))
        point[2][marker_id] = distance
    marker_ids = sorted(marker_ids)
    comparison_rows = sorted(points.values())

    write_csv(output / "raw_data.csv", RAW_HEADERS, raw_rows)
    write_csv(
        output / "movement_comparison.csv",
        ["เวลาสะสม (วินาที)", "วันเวลาที่บันทึก",
         *[f"เป้า ID {marker_id} (มม.)" for marker_id in marker_ids]],
        [[elapsed, recorded_at, *[distances.get(marker_id) for marker_id in marker_ids]]
         for elapsed, recorded_at, distances in comparison_rows],
    )
    write_excel(output / "movement_report.xlsx", comparison_rows, marker_ids, rows)


def main():
    parser = argparse.ArgumentParser(description="แก้ขนาด ArUco marker ของ session ที่ระบุ")
    parser.add_argument("--database", default="movement_data.db")
    parser.add_argument("--sessions-directory", default="sessions")
    parser.add_argument("--session-ids", nargs="+", required=True)
    parser.add_argument("--old-marker-length", type=float, default=0.05)
    parser.add_argument("--new-marker-length", type=float, default=0.032)
    parser.add_argument("--apply", action="store_true", help="ยืนยันให้แก้ไขฐานข้อมูล")
    args = parser.parse_args()
    if not args.apply:
        parser.error("คำสั่งนี้แก้ข้อมูล ใช้ --apply เพื่อยืนยัน")

    factor = args.new_marker_length / args.old_marker_length
    database_path = Path(args.database)
    backup_path = database_path.with_name(
        f"{database_path.stem}_before_marker_size_fix_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}{database_path.suffix}"
    )
    shutil.copy2(database_path, backup_path)
    connection = sqlite3.connect(database_path)
    placeholders = ",".join("?" for _ in args.session_ids)
    count = connection.execute(
        f"SELECT COUNT(*) FROM measurements WHERE session_id IN ({placeholders})", args.session_ids,
    ).fetchone()[0]
    if count == 0:
        connection.close()
        raise ValueError("ไม่พบข้อมูลใน session ที่ระบุ")
    with connection:
        connection.execute(
            f"""
            UPDATE measurements SET
                x_m = x_m * ?, y_m = y_m * ?, z_m = z_m * ?,
                dx_mm = dx_mm * ?, dy_mm = dy_mm * ?, dz_mm = dz_mm * ?,
                distance_mm = distance_mm * ?, speed_mm_s = speed_mm_s * ?
            WHERE session_id IN ({placeholders})
            """,
            [factor] * 8 + args.session_ids,
        )
    for session_id in args.session_ids:
        regenerate_reports(connection, session_id, Path(args.sessions_directory))
        print(f"Corrected and regenerated: {session_id}")
    connection.close()
    print(f"Corrected {count} rows with factor {factor:.6f}")
    print(f"Database backup: {backup_path}")


if __name__ == "__main__":
    main()
